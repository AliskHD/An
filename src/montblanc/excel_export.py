"""Excel-Export fuer den Tagesplan."""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Tagesplan

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def export_tagesplan(plan: Tagesplan, pfad: Path) -> Path:
    wb = Workbook()
    _sheet_uebersicht(wb.active, plan)
    _sheet_pro_mitarbeiter(wb.create_sheet("Pro Mitarbeiter"), plan)
    _sheet_hinweise(wb.create_sheet("Hinweise"), plan)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
    return pfad


def _sheet_uebersicht(ws, plan: Tagesplan) -> None:
    ws.title = "Tagesplan"
    ws["A1"] = f"Tagesplan Mont Blanc - {plan.datum.isoformat()}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")

    headers = [
        "Mitarbeiter", "Los-Nr.", "Artikel", "Schritt",
        "Menge", "Dauer (min)", "Start", "Ende",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    row = 4
    for z in plan.zuweisungen:
        ws.cell(row=row, column=1, value=z.mitarbeiter)
        ws.cell(row=row, column=2, value=z.los_nummer)
        ws.cell(row=row, column=3, value=z.artikel)
        ws.cell(row=row, column=4, value=z.schritt)
        ws.cell(row=row, column=5, value=z.menge)
        ws.cell(row=row, column=6, value=z.dauer_min)
        ws.cell(row=row, column=7, value=z.startzeit.strftime("%H:%M"))
        ws.cell(row=row, column=8, value=z.endzeit.strftime("%H:%M"))
        row += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def _sheet_pro_mitarbeiter(ws, plan: Tagesplan) -> None:
    ws.append(["Mitarbeiter", "Start", "Ende", "Los", "Schritt", "Menge"])
    for c in ws[1]:
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL

    gruppen: dict[str, list] = defaultdict(list)
    for z in plan.zuweisungen:
        gruppen[z.mitarbeiter].append(z)

    for name in sorted(gruppen):
        for z in sorted(gruppen[name], key=lambda x: x.startzeit):
            ws.append([
                name,
                z.startzeit.strftime("%H:%M"),
                z.endzeit.strftime("%H:%M"),
                z.los_nummer,
                z.schritt,
                z.menge,
            ])

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 16


def _sheet_hinweise(ws, plan: Tagesplan) -> None:
    ws.append(["Hinweis"])
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    if plan.nicht_eingeplante_lose:
        ws.append(["Nicht eingeplante Lose:"])
        for n in plan.nicht_eingeplante_lose:
            ws.append([n])
    if plan.freie_mitarbeiter:
        ws.append([""])
        ws.append(["Mitarbeiter ohne Zuweisung:"])
        for m in plan.freie_mitarbeiter:
            ws.append([m])
    ws.column_dimensions["A"].width = 80
